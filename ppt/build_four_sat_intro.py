#!/usr/bin/env python
"""Build a four-slide intro: classical NP-complete 4-SAT → Hamiltonian.

Renders every formula with pdflatex, embeds the PNGs, then deletes them.
Output: ppt/four_sat_intro.pptx

Facts:
  - Dutta et al., arXiv:2501.11735 (hybrid qubit–qumode knapsack / ECD-VQE)
  - Lu et al., Nat. Comput. Sci. 6, 882–893 (2026), s43588-026-01007-8
  - Hamiltonians/four_sat/four_sat_000.npz
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
HAM_NPZ = ROOT / "Hamiltonians" / "four_sat" / "four_sat_000.npz"
HAM_MANIFEST = ROOT / "Hamiltonians" / "four_sat" / "four_sat_manifest.json"
OUT = HERE / "four_sat_intro.pptx"

os.environ["PATH"] = "/Library/TeX/texbin:" + os.environ.get("PATH", "")

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x33, 0x33, 0x33)
CARD = RGBColor(0xF4, 0xF4, 0xF4)
RULE = RGBColor(0xD0, 0xD0, 0xD0)
FONT = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)
NATURE_URL = "https://www.nature.com/articles/s43588-026-01007-8"
DUTTA_ARXIV = "https://arxiv.org/abs/2501.11735"

DPI = 300
QL_MAX_PX = 3600


# ---------------------------------------------------------------------------
# Instance
# ---------------------------------------------------------------------------
def load_instance() -> dict:
    data = np.load(HAM_NPZ)
    clauses = np.asarray(data["clauses"], dtype=np.int64)
    polarities = np.asarray(data["polarities"], dtype=np.int64)
    manifest = json.loads(HAM_MANIFEST.read_text(encoding="utf-8"))
    rec = next(item for item in manifest if item["file"] == HAM_NPZ.name)
    return {
        "clauses": clauses,
        "polarities": polarities,
        "n_spins": int(np.asarray(data["num_spins"]).reshape(-1)[0]),
        "n_clauses": int(len(clauses)),
        "ground_bitstring": str(rec["ground_bitstring"]),
        "n_pauli_terms": int(rec["n_pauli_terms"]),
        "identity": float(rec["identity"]),
        "gap": int(rec["gap"]),
    }


def lit_tex(var: int, pol: int) -> str:
    i = int(var) + 1
    return rf"x_{{{i}}}" if int(pol) > 0 else rf"\bar{{x}}_{{{i}}}"


def clause_tex(variables, polarities, *, name: str | None = None) -> str:
    body = r" \vee ".join(
        lit_tex(v, p) for v, p in zip(variables, polarities, strict=True)
    )
    inner = rf"({body})"
    if name is None:
        return inner
    return rf"{name} = {inner}"


def penalty_factors(variables, polarities) -> str:
    parts = []
    for v, p in zip(variables, polarities, strict=True):
        i = int(v) + 1
        parts.append(rf"(1-x_{{{i}}})" if int(p) > 0 else rf"x_{{{i}}}")
    return "".join(parts)


def projector_factors(variables, polarities) -> str:
    parts = []
    for v, p in zip(variables, polarities, strict=True):
        i = int(v) + 1
        sign = "+" if int(p) > 0 else "-"
        parts.append(rf"\tfrac{{I {sign} Z_{{{i}}}}}{{2}}")
    return "".join(parts)


# ---------------------------------------------------------------------------
# LaTeX → PNG (temp files only)
# ---------------------------------------------------------------------------
@dataclass
class EqFig:
    path: Path
    native_w: float  # inches at which glyphs are design size


def _tex_document(
    body: str,
    *,
    wrap_align: bool,
    fontsize: int,
    baselineskip: int,
    paperwidth: float,
    paperheight: float,
) -> str:
    inner = (
        f"\\begin{{align*}}\n{body.strip()}\n\\end{{align*}}"
        if wrap_align
        else body.strip()
    )
    return rf"""\documentclass[12pt]{{article}}
\usepackage{{fix-cm}}
\usepackage{{amsmath}}
\usepackage{{amssymb}}
\usepackage{{array}}
\usepackage[
  paperwidth={paperwidth:.3f}in,
  paperheight={paperheight:.3f}in,
  margin=0.16in,
  headheight=0pt,
  headsep=0pt,
  footskip=0pt
]{{geometry}}
\pagestyle{{empty}}
\setlength{{\abovedisplayskip}}{{2pt}}
\setlength{{\belowdisplayskip}}{{2pt}}
\setlength{{\abovedisplayshortskip}}{{1pt}}
\setlength{{\belowdisplayshortskip}}{{1pt}}
\begin{{document}}
\fontsize{{{fontsize}}}{{{baselineskip}}}\selectfont
{inner}
\end{{document}}
"""


def _crop_white(path: Path, pad: int = 14) -> Image.Image:
    img = Image.open(path).convert("RGB")
    arr = np.asarray(img.convert("L"))
    ink = arr < 248
    if not ink.any():
        raise RuntimeError(f"rasterized page is blank: {path}")
    ys, xs = np.where(ink)
    box = (
        max(0, int(xs.min()) - pad),
        max(0, int(ys.min()) - pad),
        min(img.width, int(xs.max()) + pad + 1),
        min(img.height, int(ys.max()) + pad + 1),
    )
    return img.crop(box)


def _pad_tight(img: Image.Image, pad: int = 10) -> Image.Image:
    """A few pixels of white around the ink, no full-page banner."""
    canvas = Image.new(
        "RGB",
        (img.width + 2 * pad, img.height + 2 * pad),
        (255, 255, 255),
    )
    canvas.paste(img, (pad, pad))
    return canvas


def render_eq(
    out_dir: Path,
    name: str,
    body: str,
    *,
    wrap_align: bool = True,
    fontsize: int = 22,
    baselineskip: int = 28,
    paperwidth: float = 12.0,
    paperheight: float = 4.8,
) -> EqFig:
    pdflatex = shutil.which("pdflatex") or "/Library/TeX/texbin/pdflatex"
    qlmanage = "/usr/bin/qlmanage"
    if not Path(pdflatex).is_file():
        raise FileNotFoundError("pdflatex not found")
    if not Path(qlmanage).is_file():
        raise FileNotFoundError("qlmanage not found")

    tex_source = _tex_document(
        body,
        wrap_align=wrap_align,
        fontsize=fontsize,
        baselineskip=baselineskip,
        paperwidth=paperwidth,
        paperheight=paperheight,
    )
    with tempfile.TemporaryDirectory(prefix=f"eq_{name}_") as tmp:
        td = Path(tmp)
        tex_path = td / f"{name}.tex"
        tex_path.write_text(tex_source, encoding="utf-8")
        proc = subprocess.run(
            [
                pdflatex,
                "-interaction=nonstopmode",
                "-halt-on-error",
                f"-output-directory={td}",
                str(tex_path),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        pdf = td / f"{name}.pdf"
        if proc.returncode != 0 or not pdf.is_file():
            raise RuntimeError(
                f"{name} pdflatex failed\n{proc.stdout[-2500:]}\n{proc.stderr[-800:]}"
            )
        subprocess.run(
            [qlmanage, "-t", "-s", str(QL_MAX_PX), "-o", str(td), str(pdf)],
            check=True,
            capture_output=True,
            text=True,
        )
        thumb = td / f"{name}.pdf.png"
        page = Image.open(thumb)
        page_w = page.width
        cropped = _crop_white(thumb)
        img = _pad_tight(cropped)
        dest = out_dir / f"{name}.png"
        img.save(dest, "PNG")
        # Place at the physical width of the ink, not the full PDF page.
        native_w = img.width * paperwidth / page_w
        print(
            f"  {name}: {img.size[0]}x{img.size[1]}  "
            f"native {native_w:.2f}in (page {paperwidth:.2f}in)"
        )
        return EqFig(path=dest, native_w=native_w)


def render_all(eq_dir: Path, inst: dict) -> dict[str, EqFig]:
    clauses = inst["clauses"]
    poles = inst["polarities"]
    n = inst["n_clauses"]
    gs = inst["ground_bitstring"]
    gs_spaced = r"\,".join(gs)
    c0, p0 = clauses[0], poles[0]
    c1, p1 = clauses[1], poles[1]

    ncols = 3
    nrows = (n + ncols - 1) // ncols
    catalog_rows = []
    for r in range(nrows):
        cells = []
        for c in range(ncols):
            idx = c * nrows + r
            if idx < n:
                cells.append(
                    "$"
                    + clause_tex(clauses[idx], poles[idx], name=rf"C_{{{idx + 1}}}")
                    + "$"
                )
            else:
                cells.append("")
        catalog_rows.append(" & ".join(cells) + r" \\[0.12em]")
    catalog_body = "\n".join(catalog_rows)

    figs: dict[str, EqFig] = {}

    figs["sat_def"] = render_eq(
        eq_dir,
        "sat_def",
        r"""
\phi(x)
  &= \bigwedge_{j=1}^{m} C_j\,,
  \qquad
  C_j
  = (\ell_{j1}\vee \ell_{j2}\vee \ell_{j3}\vee \ell_{j4})\,,
  \qquad
  \ell\in\{x_i,\bar x_i\} \\[0.35em]
\text{4-SAT}
  &\colon
  \text{does there exist }x\in\{0,1\}^n\text{ with }\phi(x)=1\,?
""",
        fontsize=20,
        baselineskip=26,
        paperwidth=12.2,
        paperheight=2.2,
    )

    figs["sat_catalog"] = render_eq(
        eq_dir,
        "sat_catalog",
        rf"""
{{\fontsize{{15}}{{20}}\selectfont
\begin{{tabular}}{{@{{}}l@{{\hspace{{1.6em}}}}l@{{\hspace{{1.6em}}}}l@{{}}}}
{catalog_body}
\end{{tabular}}}}
""",
        wrap_align=False,
        fontsize=15,
        baselineskip=20,
        paperwidth=12.2,
        paperheight=3.6,
    )

    figs["sat_gs"] = render_eq(
        eq_dir,
        "sat_gs",
        rf"""
x^\star
  &= {gs_spaced}
  \quad(x_1 x_2\cdots x_7),
  \qquad
  \phi(x^\star)=1
  \text{{ and this assignment is unique.}}
""",
        fontsize=18,
        baselineskip=24,
        paperwidth=12.2,
        paperheight=1.4,
    )

    figs["scaling"] = render_eq(
        eq_dir,
        "scaling",
        r"""
T_{\mathrm{QAOA}}
  &\propto 1.0077^{\,n}
  \qquad
  (40\text{ layers}) \\[0.25em]
T_{\mathrm{QAA}}
  &\propto 1.0070^{\,n}
  \qquad
  (150\text{ layers}) \\[0.25em]
T_{\mathrm{classical}}
  &\propto 1.0128^{\,n}
  \qquad
  \text{SOTA SAT solvers}
""",
        fontsize=20,
        baselineskip=26,
        paperwidth=6.4,
        paperheight=3.0,
    )

    figs["map_steps"] = render_eq(
        eq_dir,
        "map_steps",
        r"""
P_C
  &= \prod_{x_i\in C}(1-x_i)\,
     \prod_{\bar x_j\in C} x_j
  \qquad\text{penalty $=1$ on the unique falsifying assignment} \\[0.38em]
H(x)
  &= \sum_{C} P_C
  \qquad\text{energy $=$ number of unsatisfied clauses} \\[0.38em]
x_i
  &= \tfrac{1-Z_i}{2}
  \qquad\text{with $Z|0\rangle=+1$ and $Z|1\rangle=-1$} \\[0.38em]
H_C
  &= \prod_{x_i\in C}\tfrac{I+Z_i}{2}\,
     \prod_{\bar x_j\in C}\tfrac{I-Z_j}{2}\,,
  \qquad
  H = \sum_{C} H_C
""",
        fontsize=18,
        baselineskip=24,
        paperwidth=12.2,
        paperheight=3.4,
    )

    figs["worked"] = render_eq(
        eq_dir,
        "worked",
        rf"""
{clause_tex(c0, p0, name=r"C_1")}
  &\qquad
  P_{{C_1}}
  = {penalty_factors(c0, p0)}
  = {projector_factors(c0, p0)} \\[0.45em]
{clause_tex(c1, p1, name=r"C_2")}
  &\qquad
  P_{{C_2}}
  = {penalty_factors(c1, p1)}
  = {projector_factors(c1, p1)} \\[0.45em]
H
  &= H_{{C_1}}+H_{{C_2}}+\cdots+H_{{C_{{{n}}}}}
  \qquad\text{{18 projectors; do not expand the Pauli sum}}
""",
        fontsize=16,
        baselineskip=22,
        paperwidth=12.2,
        paperheight=2.6,
    )

    return figs


# ---------------------------------------------------------------------------
# Slide helpers
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    bar = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SW, Inches(0.07)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLACK
    bar.line.fill.background()
    return s


def _run(p, text, size, bold=False, color=None, italic=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = BLACK if color is None else color
    return r


def _para(
    tf,
    text,
    size=16,
    bold=False,
    first=False,
    after=6,
    before=0,
    align=None,
    spacing=1.12,
    color=None,
    italic=False,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if first and tf.paragraphs[0].runs:
        p = tf.add_paragraph()
    if isinstance(text, str):
        _run(p, text, size, bold=bold, color=color, italic=italic)
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    p.line_spacing = spacing
    if align is not None:
        p.alignment = align
    return p


def _para_runs(tf, segs, size=16, first=False, after=6, before=0, spacing=1.12):
    """segs: list of (text, bold, italic, color_or_None, hyperlink_or_None)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if first and tf.paragraphs[0].runs:
        p = tf.add_paragraph()
    for rec in segs:
        text, bold, italic, color, link = rec
        r = _run(p, text, size, bold=bold, color=color, italic=italic)
        if link:
            r.hyperlink.address = link
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    p.line_spacing = spacing
    return p


def tb(s, x, y, w, h, *, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    return tf


def title(s, text):
    tf = tb(s, 0.50, 0.16, 12.4, 0.52)
    _para(tf, text, size=24, bold=True, first=True, after=0)
    line = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.50), Inches(0.70), Inches(12.33), Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLACK
    line.line.fill.background()


def card(s, x, y, w, h):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = RULE
    shp.line.width = Emu(6350)
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    return shp


def fig_h(fig: EqFig, w: float) -> float:
    im = Image.open(fig.path)
    return w * im.size[1] / im.size[0]


def pic(s, fig: EqFig, x, y, w=None):
    width = fig.native_w if w is None else w
    return s.shapes.add_picture(
        str(fig.path), Inches(x), Inches(y), width=Inches(width)
    )


def place_fig(s, fig: EqFig, x, y, max_w, gap=0.08, fill=False) -> float:
    w = float(max_w) if fill else min(float(fig.native_w), float(max_w))
    h = fig_h(fig, w)
    pic(s, fig, x, y, w=w)
    return y + h + gap


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def build_slides(figs: dict[str, EqFig], inst: dict) -> None:
    n = inst["n_clauses"]
    n_spins = inst["n_spins"]
    gs = inst["ground_bitstring"]
    n_pauli = inst["n_pauli_terms"]

    # ===== 1. Classical NP-complete, quantum still useful =====
    s = slide()
    title(s, "The problems themselves are classical")
    tf = tb(s, 0.50, 0.78, 12.3, 0.42)
    _para(
        tf,
        "0-1 knapsack and 4-SAT are ordinary combinatorial questions — and NP-complete. A quantum machine can still be the right solver.",
        size=16,
        first=True,
        after=0,
        color=MUTED,
    )

    card(s, 0.50, 1.28, 3.95, 3.22)
    tf = tb(s, 0.66, 1.40, 3.63, 2.98)
    _para(tf, "Two textbook problems", size=16, bold=True, first=True, after=8)
    _para(tf, "0-1 knapsack.", size=14, bold=True, after=2)
    _para(
        tf,
        "Items with weights and values, a capacity. Can you pack value at least V without overflowing? (Or: pack as much value as possible.)",
        size=13,
        after=8,
        spacing=1.08,
    )
    _para(tf, "4-SAT.", size=14, bold=True, after=2)
    _para(
        tf,
        "A Boolean formula whose every clause is an OR of four literals. Does any 0/1 assignment make every clause true?",
        size=13,
        after=0,
        spacing=1.08,
    )

    card(s, 4.63, 1.28, 3.95, 3.22)
    tf = tb(s, 4.79, 1.40, 3.63, 2.98)
    _para(tf, "NP-complete, in plain words", size=16, bold=True, first=True, after=8)
    _para(
        tf,
        "NP: a proposed answer can be checked quickly (polynomial time).",
        size=13,
        after=7,
        spacing=1.08,
    )
    _para(
        tf,
        "NP-complete: the hardest problems in NP. Every other NP problem can be rewritten as one of them in polynomial time.",
        size=13,
        after=7,
        spacing=1.08,
    )
    _para(
        tf,
        "Easy to verify, believed hard to find. A fast algorithm for any one of them would give a fast algorithm for all of NP.",
        size=13,
        after=0,
        spacing=1.08,
    )

    card(s, 8.76, 1.28, 4.07, 3.22)
    tf = tb(s, 8.92, 1.40, 3.75, 2.98)
    _para(tf, "Why a quantum machine can still help", size=16, bold=True, first=True, after=8)
    _para(
        tf,
        "NP-completeness is a classical worst-case label. It does not forbid better typical-case scaling, or a more compact encoding on quantum hardware.",
        size=13,
        after=7,
        spacing=1.08,
    )
    _para(
        tf,
        "Turn the cost into a Hamiltonian. Variational and adiabatic algorithms search that landscape in superposition.",
        size=13,
        after=7,
        spacing=1.08,
    )
    _para(
        tf,
        "For SAT-type problems, that is already showing empirical scaling gains — next slide.",
        size=13,
        after=0,
        spacing=1.08,
    )

    card(s, 0.50, 4.64, 12.33, 2.48)
    tf = tb(s, 0.70, 4.76, 12.0, 2.20)
    _para(tf, "Previous work chose knapsack. We choose 4-SAT.", size=16, bold=True, first=True, after=8)
    _para_runs(
        tf,
        [
            ("Previous work.  ", True, False, None, None),
            (
                "Dutta, Allen, Xu, et al., “Solving Constrained Optimization Problems Using Hybrid Qubit-Qumode Quantum Devices,” arXiv:2501.11735. ",
                False,
                False,
                None,
                None,
            ),
            ("Binary knapsack", True, False, None, None),
            (
                " as a QUBO, solved with ECD-VQE on a hybrid qubit–qumode device.  ",
                False,
                False,
                None,
                None,
            ),
            (DUTTA_ARXIV, False, False, MUTED, DUTTA_ARXIV),
        ],
        size=14,
        after=8,
        spacing=1.12,
    )
    _para_runs(
        tf,
        [
            ("This work.  ", True, False, None, None),
            (
                "Same motivation — a classical NP-complete cost, encoded as a Hamiltonian, aimed at the same hybrid architecture — but the target is ",
                False,
                False,
                None,
                None,
            ),
            ("4-SAT", True, False, None, None),
            (".", False, False, None, None),
        ],
        size=14,
        after=0,
        spacing=1.12,
    )

    # ===== 2. Nature paper as support =====
    s = slide()
    title(s, "Empirical support: quantum SAT solvers can still win")
    tf = tb(s, 0.50, 0.78, 12.3, 0.58)
    _para_runs(
        tf,
        [
            (
                "Lu, Wei, Li, Gao, Yan, Zheng, Zhang, Zeng, Long, et al.  ",
                True,
                False,
                None,
                None,
            ),
        ],
        size=14,
        first=True,
        after=1,
        spacing=1.05,
    )
    _para_runs(
        tf,
        [
            (
                "Evidence of scaling advantage on an NP-complete problem with enhanced quantum solvers.  ",
                False,
                True,
                None,
                None,
            ),
            ("Nat. Comput. Sci. 6, 882–893 (2026).", False, False, None, None),
        ],
        size=14,
        after=0,
        spacing=1.05,
    )

    card(s, 0.50, 1.42, 12.33, 1.48)
    tf = tb(s, 0.70, 1.50, 12.0, 1.32)
    _para(
        tf,
        "“Without complexity proofs, scaling advantage — where quantum resource requirements grow more slowly than their classical counterparts — is the primary indicator.”",
        size=15,
        italic=True,
        first=True,
        after=6,
        spacing=1.12,
    )
    _para(
        tf,
        "They report empirical evidence of quantum speedup on an NP-complete SAT problem. That is the argument that a quantum machine can still be beneficial here.",
        size=14,
        after=0,
        spacing=1.08,
    )

    card(s, 0.50, 3.04, 6.05, 3.72)
    tf = tb(s, 0.66, 3.14, 5.73, 3.50)
    _para(tf, "What they did", size=16, bold=True, first=True, after=7)
    _para(
        tf,
        "Target: one-in-three SAT, an NP-complete cousin of 3-SAT. Each clause is true iff exactly one of its three literals is true.",
        size=13,
        after=6,
        spacing=1.08,
    )
    _para(
        tf,
        "They built enhanced QAOA and QAA solvers, with a classical space-reduction step, and compared scaling to state-of-the-art classical SAT solvers.",
        size=13,
        after=6,
        spacing=1.08,
    )
    _para(
        tf,
        "A 13-qubit superconducting experiment solved instances with 22 variables and matched the predicted improvement.",
        size=13,
        after=6,
        spacing=1.08,
    )
    _para(
        tf,
        "Different SAT variant than our 4-SAT. Same moral: NP-complete SAT is a live target for new quantum algorithms.",
        size=13,
        after=0,
        spacing=1.08,
    )

    card(s, 6.73, 3.04, 6.10, 3.72)
    tf = tb(s, 6.89, 3.14, 5.78, 0.62)
    _para(tf, "Scaling at the critical ratio m/n = 0.626, n ≤ 70", size=15, bold=True, first=True, after=4)
    _para(
        tf,
        "Inverse success probability (quantum) vs. solver work (classical).",
        size=13,
        after=0,
        color=MUTED,
    )
    scale_y = place_fig(s, figs["scaling"], 6.89, 3.84, 5.70, gap=0.08, fill=True)
    tf = tb(s, 6.89, min(scale_y, 5.85), 5.78, 0.78)
    _para(
        tf,
        "Both enhanced quantum solvers beat 1.0128ⁿ from the best classical baselines they report.",
        size=13,
        first=True,
        after=0,
        spacing=1.08,
    )

    tf = tb(s, 0.50, 6.90, 12.3, 0.42)
    _para_runs(
        tf,
        [
            (NATURE_URL, False, False, MUTED, NATURE_URL),
        ],
        size=12,
        first=True,
        after=0,
    )

    # ===== 3. What 4-SAT is =====
    s = slide()
    title(s, "What 4-SAT is")
    tf = tb(s, 0.50, 0.78, 12.3, 0.72)
    _para(
        tf,
        "A 4-SAT instance is a Boolean formula in conjunctive normal form in which every clause has exactly four literals. A literal is a variable or its negation. The formula is satisfiable if some assignment in {0,1}ⁿ makes every clause true.",
        size=15,
        first=True,
        after=4,
        spacing=1.10,
    )
    _para(
        tf,
        "SAT was the first NP-complete problem (Cook–Levin). For every k ≥ 3, k-SAT stays NP-complete — 4-SAT included.",
        size=15,
        after=0,
        spacing=1.10,
    )

    y = place_fig(s, figs["sat_def"], 0.45, 1.58, 12.2, gap=0.10)

    tf = tb(s, 0.50, y, 12.3, 0.40)
    _para(
        tf,
        f"Running example: Hamiltonians/four_sat/four_sat_000  —  {n_spins} variables, {n} clauses, unique satisfying assignment.",
        size=15,
        bold=True,
        first=True,
        after=0,
    )
    y2 = place_fig(s, figs["sat_catalog"], 0.45, y + 0.42, 12.2, gap=0.08, fill=True)
    place_fig(s, figs["sat_gs"], 0.45, y2, 12.2, gap=0.04, fill=True)

    # ===== 4. Clause → penalty → Hamiltonian =====
    s = slide()
    title(s, "From four_sat_000 to a spin Hamiltonian")
    tf = tb(s, 0.50, 0.78, 12.3, 0.48)
    _para(
        tf,
        "Three steps. Rewrite each clause as a 0/1 penalty, add the penalties, then replace every bit by a Pauli Z. We never write the fully expanded Hamiltonian.",
        size=15,
        first=True,
        after=0,
        spacing=1.10,
    )

    y = place_fig(s, figs["map_steps"], 0.40, 1.26, 12.2, gap=0.08)

    tf = tb(s, 0.50, y, 12.3, 0.32)
    _para(
        tf,
        "Worked example — the first two clauses of four_sat_000. Positive literal → (1 − x) → (I + Z)/2. Negated literal → x → (I − Z)/2.",
        size=14,
        bold=True,
        first=True,
        after=0,
    )
    y3 = place_fig(s, figs["worked"], 0.40, y + 0.30, 12.2, gap=0.06, fill=True)

    tf = tb(s, 0.50, min(y3, 6.88), 12.3, 0.48)
    _para(
        tf,
        f"Unique ground state of four_sat_000: {gs}, energy 0, gap 1. Full Pauli expansion: {n_pauli} Z-strings plus an identity shift — stored in the NPZ, not shown.",
        size=14,
        first=True,
        after=0,
        spacing=1.08,
    )


def main() -> None:
    if not HAM_NPZ.is_file():
        raise FileNotFoundError(HAM_NPZ)
    inst = load_instance()
    print(
        f"four_sat_000: {inst['n_spins']} spins, {inst['n_clauses']} clauses, "
        f"GS={inst['ground_bitstring']}"
    )
    with tempfile.TemporaryDirectory(prefix="four_sat_ppt_eqs_") as tmp:
        eq_dir = Path(tmp)
        print("Rendering equations with pdflatex …")
        figs = render_all(eq_dir, inst)
        build_slides(figs, inst)
        prs.save(str(OUT))
        print(f"Wrote {OUT}")
    leftover = list(HERE.glob("*_eqs*.png")) + list(HERE.glob("eq_*.png"))
    if leftover:
        raise RuntimeError(f"equation PNGs were not cleaned up: {leftover}")
    print("Temporary equation images deleted.")


if __name__ == "__main__":
    main()
